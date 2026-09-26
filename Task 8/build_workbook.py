"""Step 2: build the base Excel workbook (data table, KPI formulas, analysis tables,
insights, cleaning log, data dictionary) with openpyxl.

Pivot tables, pivot charts, slicers and the Dashboard sheet are added in step 3
(build_dashboard.py) through Excel itself, because openpyxl cannot create them.
"""
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = "NoShows_Base.xlsx"
NAVY, TEAL, CORAL, GREY = "1F3B57", "1B998B", "E4572E", "F3F6F9"
FONT = "Arial"

df = pd.read_csv("Cleaned_NoShows.csv", dtype={"Patient_ID": str}, keep_default_na=False,
                 parse_dates=["Scheduled_Date", "Appointment_Date"])
log = json.load(open("cleaning_log.json"))
N = len(df)
LAST = N + 1

wb = Workbook()
thin = Side(style="thin", color="D0D7DE")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def hdr(ws, row, values, col=1, fill=NAVY):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, italic=True, size=10, color="5A6B7B")
    ws.sheet_view.showGridLines = False


def body(c, fmt=None, bold=False, color="000000"):
    c.font = Font(name=FONT, size=10, bold=bold, color=color)
    c.border = BOX
    if fmt:
        c.number_format = fmt


# ---------------------------------------------------------------- Data sheet
ws = wb.active
ws.title = "Data"
cols = list(df.columns)
ws.append(cols)
for row in df.itertuples(index=False):
    ws.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in row])
L = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
for c in ("Scheduled_Date", "Appointment_Date"):
    for cell in ws[L[c]][1:]:
        cell.number_format = "yyyy-mm-dd"
tab = Table(displayName="tblAppointments", ref=f"A1:{get_column_letter(len(cols))}{LAST}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tab)
ws.freeze_panes = "C2"
for i, c in enumerate(cols):
    ws.column_dimensions[get_column_letter(i + 1)].width = max(11, min(24, len(c) + 3))


def R(col):
    """Absolute range of a Data column (fast in formulas, stable in every Excel version)."""
    return f"Data!${L[col]}$2:${L[col]}${LAST}"


# ---------------------------------------------------------------- KPIs sheet
k = wb.create_sheet("KPIs")
title(k, "Key Performance Indicators (all appointments)",
      "Live formulas on the Data sheet. The Dashboard KPI cards show the same measures for the slicer selection.")
hdr(k, 4, ["#", "KPI", "Value", "Formula logic", "Why it matters"])
NS, ST, LD, FA = R("No_Show"), R("Status"), R("Lead_Days"), R("First_Appt")
SMS, HIST, SCH, AGE = R("SMS_Received"), R("NoShow_History"), R("Scholarship"), R("Age")
kpis = [
    ("Total appointments", f"=COUNTA({R('Appointment_ID')})", "#,##0", "COUNTA of Appointment_ID",
     "Workload / demand volume"),
    ("Unique patients", f"=SUM({FA})", "#,##0", "SUM of First_Appt (1 on each patient's first row)",
     "Size of the patient base"),
    ("Appointments per patient", "=C5/C6", "0.00", "Appointments / patients", "How often patients return"),
    ("Showed up", f'=COUNTIF({ST},"Showed Up")', "#,##0", 'COUNTIF Status = "Showed Up"', "Completed visits"),
    ("No-shows", f"=SUM({NS})", "#,##0", "SUM of No_Show (1 = missed)", "Wasted slots"),
    ("No-show rate", "=C9/C5", "0.0%", "No-shows / appointments", "Headline KPI: share of slots lost"),
    ("Attendance rate", "=C8/C5", "0.0%", "Showed up / appointments", "Complement of no-show rate"),
    ("Avg lead time (days)", f"=AVERAGE({LD})", "0.0", "AVERAGE of Lead_Days", "Booking-to-visit wait"),
    ("Avg lead time - no-shows", f'=AVERAGEIF({ST},"No-Show",{LD})', "0.0", "AVERAGEIF Status = No-Show",
     "Compare with the next KPI"),
    ("Avg lead time - showed up", f'=AVERAGEIF({ST},"Showed Up",{LD})', "0.0", "AVERAGEIF Status = Showed Up",
     "No-shows waited ~2x longer"),
    ("Same-day bookings share", f"=COUNTIF({LD},0)/C5", "0.0%", "Lead_Days = 0 / all", "Walk-in style demand"),
    ("No-show rate - same day", f"=SUMIFS({NS},{LD},0)/COUNTIF({LD},0)", "0.0%", "Rate where Lead_Days = 0",
     "Baseline: almost everyone attends"),
    ("No-show rate - booked 1+ days ahead", f'=SUMIFS({NS},{LD},">0")/COUNTIF({LD},">0")', "0.0%",
     "Rate where Lead_Days > 0", "The real risk population"),
    ("No-show rate - booked 15+ days ahead", f'=SUMIFS({NS},{LD},">=15")/COUNTIF({LD},">=15")', "0.0%",
     "Rate where Lead_Days >= 15", "Long waits drive no-shows"),
    ("Share of no-shows from 15+ day bookings", f'=SUMIFS({NS},{LD},">=15")/C9', "0.0%",
     "No-shows with Lead_Days >= 15 / all no-shows", "Where to target reminders"),
    ("SMS coverage (1+ day bookings)", f'=COUNTIFS({SMS},"Yes",{LD},">0")/COUNTIF({LD},">0")', "0.0%",
     "SMS sent / bookings with Lead_Days > 0", "No SMS is ever sent for same-day bookings"),
    ("No-show rate 1+ days - SMS", f'=SUMIFS({NS},{SMS},"Yes",{LD},">0")/COUNTIFS({SMS},"Yes",{LD},">0")',
     "0.0%", "Rate, SMS = Yes, Lead_Days > 0", "Fair SMS comparison (same lead-time population)"),
    ("No-show rate 1+ days - no SMS", f'=SUMIFS({NS},{SMS},"No",{LD},">0")/COUNTIFS({SMS},"No",{LD},">0")',
     "0.0%", "Rate, SMS = No, Lead_Days > 0", "SMS lowers no-shows once lead time is controlled"),
    ("No-show rate - patients who missed before",
     f'=SUMIFS({NS},{HIST},"Has missed before")/COUNTIF({HIST},"Has missed before")', "0.0%",
     "Rate where earlier no-show exists", "Past behaviour predicts future no-show"),
    ("No-show rate - patients who never missed",
     f'=SUMIFS({NS},{HIST},"Never missed")/COUNTIF({HIST},"Never missed")', "0.0%",
     "Rate for returning patients with a clean record", "Reliable patients"),
    ("No-show rate - scholarship (Bolsa Familia)", f'=SUMIFS({NS},{SCH},"Yes")/COUNTIF({SCH},"Yes")', "0.0%",
     "Rate where Scholarship = Yes", "Socio-economic barrier"),
    ("No-show rate - ages 13-30", f'=SUMIFS({NS},{AGE},">=13",{AGE},"<=30")/COUNTIFS({AGE},">=13",{AGE},"<=30")',
     "0.0%", "Rate for teens and young adults", "Highest-risk age band"),
    ("No-show rate - ages 61+", f'=SUMIFS({NS},{AGE},">=61")/COUNTIF({AGE},">=61")', "0.0%",
     "Rate for seniors", "Lowest-risk age band"),
]
for i, (name, f, fmt, logic, why) in enumerate(kpis):
    r = 5 + i
    vals = [i + 1, name, f, logic, why]
    for j, v in enumerate(vals):
        c = k.cell(r, 1 + j, v)
        body(c, fmt if j == 2 else None, bold=(j == 2), color=NAVY if j == 2 else "000000")
        if i % 2:
            c.fill = PatternFill("solid", fgColor=GREY)
for col, w in zip("ABCDE", (5, 40, 14, 48, 48)):
    k.column_dimensions[col].width = w
k.freeze_panes = "A5"

# ----------------------------------------------------------- Analysis sheet
a = wb.create_sheet("Analysis")
title(a, "No-show analysis by factor (formula tables)",
      "Every number is a COUNTIFS/SUMIFS formula on the Data sheet. Index = factor rate / overall rate "
      "(>1 = riskier than average).")
row = 4


def block(field, heading, values=None, filt=None):
    """Write a factor table: category | appointments | no-shows | rate | share of no-shows | index."""
    global row
    a.cell(row, 1, heading).font = Font(name=FONT, bold=True, size=12, color=NAVY)
    row += 1
    hdr(a, row, [field, "Appointments", "No-shows", "No-show rate", "Share of no-shows", "Risk index"])
    row += 1
    top = row
    cats = values if values is not None else sorted(df[field].unique())
    for v in cats:
        a.cell(row, 1, v)
        a.cell(row, 2, f'=COUNTIF({R(field)},A{row})')
        a.cell(row, 3, f'=SUMIFS({NS},{R(field)},A{row})')
        a.cell(row, 4, f"=IFERROR(C{row}/B{row},0)")
        a.cell(row, 5, f"=C{row}/KPIs!$C$9")
        a.cell(row, 6, f"=D{row}/KPIs!$C$10")
        for j, fmt in enumerate([None, "#,##0", "#,##0", "0.0%", "0.0%", "0.00"]):
            body(a.cell(row, 1 + j), fmt)
        row += 1
    a.conditional_formatting.add(f"D{top}:D{row - 1}", ColorScaleRule(
        start_type="min", start_color="E8F6F3", mid_type="percentile", mid_value=50, mid_color="FFF4E0",
        end_type="max", end_color="F4B6A6"))
    a.conditional_formatting.add(f"B{top}:B{row - 1}", DataBarRule(start_type="min", end_type="max", color="9BC4E2"))
    row += 2


block("Lead_Time_Group", "1. Lead time (days between booking and appointment)")
block("Age_Group", "2. Age group")
block("NoShow_History", "3. Patient no-show history", ["First visit", "Never missed", "Has missed before"])
block("SMS_Received", "4. SMS reminder (all bookings - misleading, see table 5)")

# SMS x lead time matrix: the Simpson's paradox table
a.cell(row, 1, "5. SMS effect controlled for lead time: no-show rate").font = Font(name=FONT, bold=True, size=12, color=NAVY)
row += 1
hdr(a, row, ["Lead_Time_Group", "Appts no SMS", "Appts with SMS", "Rate no SMS", "Rate with SMS", "SMS lift (pts)"])
row += 1
top = row
for v in sorted(df.Lead_Time_Group.unique()):
    lt = R("Lead_Time_Group")
    a.cell(row, 1, v)
    a.cell(row, 2, f'=COUNTIFS({lt},A{row},{SMS},"No")')
    a.cell(row, 3, f'=COUNTIFS({lt},A{row},{SMS},"Yes")')
    a.cell(row, 4, f'=IFERROR(SUMIFS({NS},{lt},A{row},{SMS},"No")/B{row},"-")')
    a.cell(row, 5, f'=IFERROR(SUMIFS({NS},{lt},A{row},{SMS},"Yes")/C{row},"-")')
    a.cell(row, 6, f'=IFERROR((E{row}-D{row})*100,"-")')
    for j, fmt in enumerate([None, "#,##0", "#,##0", "0.0%", "0.0%", "+0.0;-0.0;0.0"]):
        body(a.cell(row, 1 + j), fmt)
    row += 1
a.cell(row, 1, "Negative lift = fewer no-shows with SMS. SMS is only sent for bookings made in advance, which are "
               "riskier, so the raw comparison in table 4 hides the benefit.").font = Font(name=FONT, italic=True, size=9, color="5A6B7B")
row += 3

block("Scholarship", "6. Scholarship (Bolsa Familia welfare programme)")
block("Chronic_Group", "7. Chronic conditions (hypertension, diabetes, alcoholism, handicap)",
      ["None", "1 condition", "2+ conditions"])
for cond in ["Hypertension", "Diabetes", "Alcoholism", "Handicap"]:
    block(cond, f"   {cond}")
block("Gender", "8. Gender")
block("Appt_Weekday", "9. Appointment weekday")
block("Patient_Type", "10. Patient type (number of appointments in the period)", ["One-time", "2-4 visits", "5+ visits"])
nb = df.groupby("Neighbourhood").No_Show.agg(["count", "mean"])
nb = nb[nb["count"] >= 500].sort_values("mean", ascending=False)
block("Neighbourhood", "11. Neighbourhoods with 500+ appointments (sorted by no-show rate, highest first)", list(nb.index))
for col, w in zip("ABCDEF", (44, 15, 15, 15, 17, 14)):
    a.column_dimensions[col].width = w

# ----------------------------------------------------------- Cleaning log
c = wb.create_sheet("Cleaning Log")
title(c, "Data cleaning and preparation log", f"Raw file: healthcare_noshows.csv (106,987 rows). Clean table: {N:,} rows.")
hdr(c, 4, ["Step", "Check", "Finding", "Action taken", "Rows after"])
for i, r in enumerate(log):
    for j, v in enumerate([r["step"], r["check"], r["found"], r["action"], r["rows_after"]]):
        cell = c.cell(5 + i, 1 + j, v)
        body(cell, "#,##0" if j == 4 else None)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDE", (6, 30, 60, 60, 12)):
    c.column_dimensions[col].width = w

# ----------------------------------------------------------- Data dictionary
d = wb.create_sheet("Data Dictionary")
title(d, "Data dictionary (tblAppointments)")
hdr(d, 3, ["Column", "Type", "Source", "Description"])
DESC = {
    "Appointment_ID": ("Integer", "Raw", "Unique appointment identifier"),
    "Patient_ID": ("Text", "Raw (cleaned)", "Patient identifier, kept as text to avoid precision loss"),
    "Gender": ("Text", "Raw (cleaned)", "Female / Male"),
    "Age": ("Integer", "Raw", "Age in years (ages of 115 flagged)"),
    "Age_Group": ("Text", "Derived", "7 life-stage bands, prefixed for sort order"),
    "Neighbourhood": ("Text", "Raw (cleaned)", "Neighbourhood of the health unit, Title Case"),
    "Scholarship": ("Yes/No", "Raw", "Enrolled in the Bolsa Familia welfare programme"),
    "Hypertension": ("Yes/No", "Raw (renamed)", "Hypertension (raw: Hipertension)"),
    "Diabetes": ("Yes/No", "Raw", "Diabetes"),
    "Alcoholism": ("Yes/No", "Raw", "Alcoholism"),
    "Handicap": ("Yes/No", "Raw (renamed)", "Handicap (raw: Handcap)"),
    "Chronic_Count": ("Integer", "Derived", "Number of the 4 conditions above"),
    "Chronic_Group": ("Text", "Derived", "None / 1 condition / 2+ conditions"),
    "SMS_Received": ("Yes/No", "Raw", "At least one SMS reminder was sent"),
    "SMS_Flag": ("1/0", "Derived", "Numeric SMS flag (average = SMS coverage)"),
    "Scheduled_Date": ("Date", "Raw", "Date the appointment was booked"),
    "Appointment_Date": ("Date", "Raw", "Date of the appointment (2016-04-29 to 2016-06-08)"),
    "Appt_Month": ("Text", "Derived", "yyyy-mm of the appointment"),
    "Appt_Weekday": ("Text", "Derived", "Weekday of the appointment, numbered for sort order"),
    "Lead_Days": ("Integer", "Raw (renamed)", "Days between booking and appointment (raw: Date.diff)"),
    "Lead_Time_Group": ("Text", "Derived", "7 lead-time bands"),
    "Status": ("Text", "Derived", "Showed Up / No-Show (from raw Showed_up)"),
    "No_Show": ("1/0", "Derived", "1 = patient did not attend. Average = no-show rate"),
    "Patient_Appt_Count": ("Integer", "Derived", "Appointments the patient has in the dataset"),
    "Patient_Type": ("Text", "Derived", "One-time / 2-4 visits / 5+ visits"),
    "Visit_Number": ("Integer", "Derived", "Chronological order of this appointment for the patient"),
    "Prior_NoShows": ("Integer", "Derived", "No-shows the patient had before this appointment"),
    "NoShow_History": ("Text", "Derived", "First visit / Never missed / Has missed before"),
    "First_Appt": ("1/0", "Derived", "1 on the patient's first appointment. Sum = unique patients"),
    "Flag_PatientID": ("Text", "Quality flag", "Malformed ID (decimal part in raw ID)"),
    "Flag_Age": ("Text", "Quality flag", "Implausible age (115)"),
    "Flag_MultiBooking": ("Text", "Quality flag", "Patient has 2+ appointments booked and held on the same days"),
}
for i, col in enumerate(cols):
    t, s, desc = DESC[col]
    for j, v in enumerate([col, t, s, desc]):
        body(d.cell(4 + i, 1 + j, v))
for col, w in zip("ABCD", (22, 10, 16, 70)):
    d.column_dimensions[col].width = w

wb.save(OUT)
print("saved", OUT, "rows", N, "KPIs", len(kpis))
