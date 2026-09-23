"""
Task 7 - Hospital Analytics: Excel deliverable with live measures.

Reads  Hospital Analytics.xlsx, Cleaned_Hospital_Data.csv, cleaning_log.json
Writes Hospital_Analytics_Clean.xlsx

Sheets
  About            - what is in the workbook and how the measures work
  KPIs             - every headline KPI as a live formula over the Clean Data sheet, with its purpose
  Summary Tables   - COUNTIFS / AVERAGEIFS / SUMIFS tables by diagnosis, doctor, insurance, severity, age, month, weekday
  Severity x Dx    - admissions matrix (colour scale) - is severity consistent with the diagnosis?
  Daily Census     - patients in a bed each night (feeds Avg / Peak census KPIs)
  Clean Data       - 247 cleaned rows as an Excel Table (tblPatients) with calculated columns
  Cleaning Log     - each check, how many rows it hit, and the action taken
  Data Dictionary  - every column: meaning, type, source (raw / calculated)
  Raw Data         - the untouched original sheet, for audit

openpyxl writes formulas without results, so the script finishes by opening the file in LibreOffice (headless),
which calculates every formula and saves the values - the workbook then shows numbers in any viewer, not only Excel.
"""

import json
import os
import shutil
import subprocess
import tempfile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = "Hospital_Analytics_Clean.xlsx"
df = pd.read_csv("Cleaned_Hospital_Data.csv", parse_dates=["Admission_Date", "Discharge_Date"])
log = json.load(open("cleaning_log.json", encoding="utf-8"))
N = len(df)
LAST = N + 1  # last data row on Clean Data

FONT = "Arial"
NAVY = "1F3A5F"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color="52514E")
F_SECTION = Font(name=FONT, size=12, bold=True, color=NAVY)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_BAND = PatternFill("solid", fgColor="EEF3F9")
FILL_TOTAL = PatternFill("solid", fgColor="DCE6F2")
THIN = Side(style="thin", color="C3C2B7")
BOX = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
PCT, DEC1, DEC2, INT, MONEY = "0.0%", "0.0", "0.00", "#,##0", "#,##0"

wb = Workbook()


# ------------------------------------------------------------------ helpers
def col_of(name):
    return get_column_letter(list(df.columns).index(name) + 1)


def rng(name):
    c = col_of(name)
    return f"'Clean Data'!${c}$2:${c}${LAST}"


def header(ws, row, labels, col=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 30


def title(ws, text, sub):
    ws["A1"], ws["A2"] = text, sub
    ws["A1"].font, ws["A2"].font = F_TITLE, F_SUB
    ws.sheet_view.showGridLines = False


def widths(ws, w):
    for i, v in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = v


def base_font(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font() or (c.font.name != FONT):
                c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic,
                              color=c.font.color)


# ------------------------------------------------------------------ Clean Data (Excel Table)
ws_clean = wb.active
ws_clean.title = "Clean Data"
ws_clean.append(list(df.columns))
for rec in df.itertuples(index=False):
    ws_clean.append([None if pd.isna(v) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else
                                              (v.item() if hasattr(v, "item") else v)) for v in rec])
tab = Table(displayName="tblPatients", ref=f"A1:{get_column_letter(df.shape[1])}{LAST}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_clean.add_table(tab)
for c in ("Admission_Date", "Discharge_Date"):
    for cell in ws_clean[col_of(c)][1:]:
        cell.number_format = "yyyy-mm-dd"
for c in ("Bill", "Bill_Raw", "Bill_Per_Day"):
    for cell in ws_clean[col_of(c)][1:]:
        cell.number_format = "#,##0"
for i, name in enumerate(df.columns, 1):
    ws_clean.column_dimensions[get_column_letter(i)].width = max(11, min(24, len(name) + 3))
ws_clean.freeze_panes = "C2"

# ------------------------------------------------------------------ Daily Census
ws_cen = wb.create_sheet("Daily Census")
title(ws_cen, "Daily census", "Patients occupying a bed each night: admitted on/before the date and discharged after it.")
header(ws_cen, 4, ["Date", "Patients in bed", "Admissions that day", "Discharges that day"])
days = pd.date_range(df["Admission_Date"].min(), df["Discharge_Date"].max() - pd.Timedelta(days=1))
adm, dis = rng("Admission_Date"), rng("Discharge_Date")
for i, d in enumerate(days, 5):
    ws_cen.cell(row=i, column=1, value=d.to_pydatetime()).number_format = "yyyy-mm-dd"
    ws_cen.cell(row=i, column=2, value=f'=COUNTIFS({adm},"<="&A{i},{dis},">"&A{i})')
    ws_cen.cell(row=i, column=3, value=f"=COUNTIFS({adm},A{i})")
    ws_cen.cell(row=i, column=4, value=f"=COUNTIFS({dis},A{i})")
CEN_LAST = 4 + len(days)
ws_cen.conditional_formatting.add(f"B5:B{CEN_LAST}", DataBarRule(start_type="num", start_value=0, end_type="max", color="2A78D6"))
widths(ws_cen, [14, 16, 20, 20])
ws_cen.freeze_panes = "A5"

# ------------------------------------------------------------------ KPIs
ws_k = wb.create_sheet("KPIs", 0)
title(ws_k, "Hospital performance - KPIs & measures",
      f"Every value is a live formula over the 'Clean Data' sheet ({N} admissions, 2023). Edit the data and they update.")
header(ws_k, 4, ["Area", "KPI / measure", "Value", "How it is calculated", "Why it matters (question it answers)"])
L, LS, HS, UN, BILL, BPD = (rng(c) for c in ("Length_of_Stay", "Long_Stay", "High_Severity", "Is_Uninsured", "Bill", "Bill_Per_Day"))
AGE, DOC, DC, RF = rng("Age"), rng("Doctor"), rng("Date_Corrected"), rng("Review_Flag")
PID = rng("Patient_ID")
nights_2023 = (f'(SUMPRODUCT(({dis}<=DATE(2024,1,1))*{dis}+({dis}>DATE(2024,1,1))*DATE(2024,1,1))'
               f'-SUMPRODUCT(({adm}>=DATE(2023,1,1))*{adm}+({adm}<DATE(2023,1,1))*DATE(2023,1,1)))')
KPIS = [
    ("Volume", "Total admissions", f"=COUNTA({PID})", INT, "Count of admission records after removing the duplicate",
     "How much demand does the hospital serve?"),
    ("Volume", "Admissions per doctor", "=C5/SUMPRODUCT(1/COUNTIF({0},{0}))".format(DOC), DEC1,
     "Admissions / number of distinct doctors", "Is the workload spread evenly across the 7 doctors?"),
    ("Volume", "Average daily census", f"={nights_2023}/365", DEC2,
     "Occupied bed-nights inside 2023 / 365", "How many beds are needed on a normal day?"),
    ("Volume", "Peak daily census", f"=MAX('Daily Census'!B5:B{CEN_LAST})", INT,
     "Highest value on the Daily Census sheet", "Capacity needed on the busiest night - size beds and staff for this."),
    ("Efficiency", "Average length of stay (ALOS, days)", f"=AVERAGE({L})", DEC2,
     "Mean of Discharge - Admission", "Core efficiency measure: lower ALOS frees beds."),
    ("Efficiency", "Median length of stay (days)", f"=MEDIAN({L})", DEC1,
     "Middle stay length", "Typical patient - the gap to ALOS shows how much long stays pull the average."),
    ("Efficiency", "Total bed-days", f"=SUM({L})", INT, "Sum of all stay lengths", "Total bed capacity consumed."),
    ("Efficiency", "Long-stay rate (> 7 days)", f"=COUNTIF({LS},TRUE)/C5", PCT,
     "Share of admissions with stay > 7 days", "How many patients become bed-blockers?"),
    ("Efficiency", "Bed-days used by 15+ day stays", f'=SUMIF({L},">=15")/C11', PCT,
     "Bed-days of stays >= 15 days / all bed-days", "How much capacity a small group of very long stays absorbs."),
    ("Acuity", "High-severity share", f"=COUNTIF({HS},TRUE)/C5", PCT,
     "Admissions marked High severity / all", "How acute is the case mix?"),
    ("Acuity", "Senior patients share (60+)", f'=COUNTIF({AGE},">=60")/C5', PCT,
     "Admissions aged 60+ / all", "Older patients stay longer - drives geriatric resource planning."),
    ("Acuity", "Average patient age", f"=AVERAGE({AGE})", DEC1, "Mean age", "Who the hospital mainly serves."),
    ("Finance", "Uninsured rate", f"=COUNTIF({UN},TRUE)/C5", PCT,
     "Admissions with no insurance / all", "Share of revenue exposed to non-payment."),
    ("Finance", "Billing completeness", f"=COUNT({BILL})/C5", PCT,
     "Admissions with a genuine bill / all (999 & 3852 placeholders excluded)", "Can revenue be reported at all? (target: 100%)"),
    ("Finance", "Average recorded bill", f"=AVERAGE({BILL})", MONEY,
     "Mean of genuine bills only", "Typical charge per admission (indicative - short stays only)."),
    ("Finance", "Median bill per bed-day", f"=MEDIAN({BPD})", MONEY,
     "Median of Bill / Length_of_Stay on genuine bills", "Unit price of a bed-day - use to estimate lost revenue."),
    ("Finance", "Total recorded billing", f"=SUM({BILL})", MONEY,
     "Sum of genuine bills", "Revenue that is actually documented."),
    ("Data quality", "Date-correction rate", f"=COUNTIF({DC},TRUE)/C5", PCT,
     "Rows whose admission/discharge dates were reversed", "Reliability of the admission system's date entry."),
    ("Data quality", "Records flagged for clinical review", f'=COUNTIF({RF},"?*")/C5', PCT,
     "Long stay for a minor condition, or atypical age for diagnosis", "Coding / care-pathway audit workload."),
]
for i, (area, name, f, fmt, how, why) in enumerate(KPIS, 5):
    ws_k.cell(row=i, column=1, value=area).font = F_BASE
    ws_k.cell(row=i, column=2, value=name).font = F_BOLD
    v = ws_k.cell(row=i, column=3, value=f)
    v.number_format, v.font = fmt, Font(name=FONT, size=11, bold=True, color=NAVY)
    ws_k.cell(row=i, column=4, value=how).font = F_BASE
    ws_k.cell(row=i, column=5, value=why).font = F_BASE
    for c in range(1, 6):
        ws_k.cell(row=i, column=c).alignment = WRAP
        ws_k.cell(row=i, column=c).border = BOX
    if area in ("Volume", "Acuity", "Data quality"):
        for c in range(1, 6):
            ws_k.cell(row=i, column=c).fill = FILL_BAND
widths(ws_k, [13, 34, 13, 48, 58])
ws_k.freeze_panes = "A5"
note_row = 5 + len(KPIS) + 1
ws_k.cell(row=note_row, column=1, value=(
    "Note: 195 of 247 bills held the placeholder values 999 or 3852 (see Cleaning Log). They are excluded, so every finance "
    "KPI describes the 52 genuine bills - almost all short stays - and understates the true average charge.")).font = F_SUB
ws_k.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=5)
ws_k.cell(row=note_row, column=1).alignment = WRAP
ws_k.row_dimensions[note_row].height = 30

# ------------------------------------------------------------------ Summary Tables
ws_s = wb.create_sheet("Summary Tables", 1)
title(ws_s, "Summary tables", "COUNTIFS / AVERAGEIFS / SUMIFS measures over the Clean Data sheet - one block per dimension.")
COLS = ["Admissions", "% of admissions", "ALOS (days)", "Bed-days", "% of bed-days", "Long-stay rate",
        "High-severity share", "Uninsured rate", "Genuine bills", "Avg genuine bill", "Clinical review flags"]
FMTS = [INT, PCT, DEC2, INT, PCT, PCT, PCT, PCT, INT, MONEY, INT]
DIMS = [
    ("Diagnosis", "Diagnosis", df["Diagnosis"].value_counts().index.tolist()),
    ("Diagnosis group", "Diagnosis_Group", df["Diagnosis_Group"].value_counts().index.tolist()),
    ("Doctor", "Doctor", df["Doctor"].value_counts().index.tolist()),
    ("Insurance", "Insurance", ["Private", "Government", "Uninsured"]),
    ("Severity", "Severity", ["Low", "Medium", "High"]),
    ("Age group", "Age_Group", ["Child (0-17)", "Adult (18-39)", "Middle age (40-59)", "Senior (60+)"]),
    ("Gender (see caveat in Cleaning Log)", "Gender", ["Female", "Male"]),
    ("Admission month", "Admission_Month_Name", ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]),
    ("Admission weekday", "Admission_Weekday", ["Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]),
    ("Stay band", "Stay_Band", ["1-3 days", "4-7 days", "8-14 days", "15+ days"]),
]
row = 4
for label, col, cats in DIMS:
    ws_s.cell(row=row, column=1, value=f"By {label.lower()}").font = F_SECTION
    row += 1
    header(ws_s, row, [label] + COLS)
    first = row + 1
    D = rng(col)
    for cat in cats:
        row += 1
        r = row
        ws_s.cell(row=r, column=1, value=cat).font = F_BOLD
        fs = [
            f'=COUNTIFS({D},$A{r})',
            f"=B{r}/COUNTA({PID})",
            f'=IFERROR(AVERAGEIFS({L},{D},$A{r}),"")',
            f'=SUMIFS({L},{D},$A{r})',
            f"=E{r}/SUM({L})",
            f'=IFERROR(COUNTIFS({D},$A{r},{LS},TRUE)/B{r},"")',
            f'=IFERROR(COUNTIFS({D},$A{r},{HS},TRUE)/B{r},"")',
            f'=IFERROR(COUNTIFS({D},$A{r},{UN},TRUE)/B{r},"")',
            f'=COUNTIFS({D},$A{r},{BILL},"<>")',
            f'=IFERROR(AVERAGEIFS({BILL},{D},$A{r}),"-")',
            f'=COUNTIFS({D},$A{r},{RF},"?*")',
        ]
        for j, (f, fmt) in enumerate(zip(fs, FMTS), 2):
            c = ws_s.cell(row=r, column=j, value=f)
            c.number_format, c.border = fmt, BOX
    last = row
    row += 1
    ws_s.cell(row=row, column=1, value="Total").font = F_BOLD
    for j, letter in [(2, "B"), (3, "C"), (5, "E"), (6, "F"), (10, "J"), (12, "L")]:
        c = ws_s.cell(row=row, column=j, value=f"=SUM({letter}{first}:{letter}{last})")
        c.number_format, c.font = FMTS[j - 2], F_BOLD
    c = ws_s.cell(row=row, column=4, value=f"=E{row}/B{row}")
    c.number_format, c.font = DEC2, F_BOLD
    for j in range(1, 13):
        ws_s.cell(row=row, column=j).fill = FILL_TOTAL
    ws_s.conditional_formatting.add(f"B{first}:B{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="2A78D6"))
    ws_s.conditional_formatting.add(f"D{first}:D{last}", ColorScaleRule(start_type="min", start_color="F0EFEC", end_type="max", end_color="EB6834"))
    row += 3
widths(ws_s, [36, 12, 12, 11, 10, 11, 11, 12, 11, 11, 12, 12])

# ------------------------------------------------------------------ Severity x Diagnosis
ws_m = wb.create_sheet("Severity x Dx", 2)
title(ws_m, "Severity mix by diagnosis",
      "Row % of each diagnosis by recorded severity. Watch for mismatches: stroke mostly 'Low', common cold mostly 'High'.")
SEV = ["Low", "Medium", "High"]
header(ws_m, 4, ["Diagnosis"] + [f"{s} (n)" for s in SEV] + ["Total"] + [f"{s} %" for s in SEV] + ["ALOS (days)"])
DX, SV = rng("Diagnosis"), rng("Severity")
dx_list = df["Diagnosis"].value_counts().index.tolist()
for i, dx in enumerate(dx_list, 5):
    ws_m.cell(row=i, column=1, value=dx).font = F_BOLD
    for j, s in enumerate(SEV, 2):
        ws_m.cell(row=i, column=j, value=f'=COUNTIFS({DX},$A{i},{SV},"{s}")')
    ws_m.cell(row=i, column=5, value=f"=SUM(B{i}:D{i})").font = F_BOLD
    for j, src in enumerate("BCD", 6):
        ws_m.cell(row=i, column=j, value=f"={src}{i}/$E{i}").number_format = "0%"
    ws_m.cell(row=i, column=9, value=f"=AVERAGEIFS({L},{DX},$A{i})").number_format = DEC2
m_last = 4 + len(dx_list)
ws_m.conditional_formatting.add(f"F5:H{m_last}", ColorScaleRule(start_type="num", start_value=0, start_color="FCFCFB",
                                                                 end_type="num", end_value=0.8, end_color="1C5CAB"))
widths(ws_m, [18, 10, 11, 10, 9, 9, 10, 9, 12])

# ------------------------------------------------------------------ Cleaning Log
ws_log = wb.create_sheet("Cleaning Log")
title(ws_log, "Cleaning log", f"{log['raw_rows']} raw rows -> {log['clean_rows']} clean rows. Produced by clean_and_validate.py.")
header(ws_log, 4, ["#", "Check", "Rows affected", "Action taken", "Evidence / detail"])
for i, c in enumerate(log["checks"], 5):
    for j, v in enumerate([i - 4, c["check"], c["count"], c["action"], c["detail"]], 1):
        cell = ws_log.cell(row=i, column=j, value=v)
        cell.alignment, cell.border = WRAP, BOX
widths(ws_log, [5, 42, 10, 44, 80])

# ------------------------------------------------------------------ Data Dictionary
DICT = {
    "Patient_ID": ("Unique admission key (P0001...) - names repeat across different people", "Calculated"),
    "Patient_Name": ("Patient name (Arabic, as recorded) - NOT unique", "Raw"),
    "Age": ("Age in years (1-90)", "Raw"),
    "Age_Group": ("Child 0-17 / Adult 18-39 / Middle age 40-59 / Senior 60+", "Calculated"),
    "Gender": ("Male / Female (recorded; disagrees with first name in 53% of rows)", "Raw, translated"),
    "Gender_AR": ("Original Arabic gender label", "Raw"),
    "Diagnosis": ("Diagnosis in English", "Raw, translated"),
    "Diagnosis_AR": ("Original Arabic diagnosis", "Raw"),
    "Diagnosis_Group": ("Respiratory infection / Chronic disease / Trauma / Neurological", "Calculated"),
    "Doctor": ("Treating doctor (transliterated)", "Raw, translated"),
    "Doctor_AR": ("Original Arabic doctor name", "Raw"),
    "Admission_Date": ("Admission date (swapped with discharge where reversed)", "Raw, corrected"),
    "Discharge_Date": ("Discharge date", "Raw, corrected"),
    "Length_of_Stay": ("Discharge - Admission, days (1-30)", "Calculated"),
    "Stay_Band": ("1-3 / 4-7 / 8-14 / 15+ days", "Calculated"),
    "Long_Stay": ("TRUE when stay > 7 days", "Calculated"),
    "LOS_Outlier": (f"TRUE when stay > {log['los_fence']:.0f} days (Q3 + 1.5 x IQR)", "Calculated"),
    "Admission_Month": ("Month number 1-12", "Calculated"),
    "Admission_Month_Name": ("Jan...Dec", "Calculated"),
    "Admission_Quarter": ("Q1-Q4", "Calculated"),
    "Admission_Weekday": ("Day of week of admission", "Calculated"),
    "Is_Weekend": ("TRUE for Friday/Saturday admissions (regional weekend)", "Calculated"),
    "Insurance": ("Private / Government / Uninsured", "Raw, translated"),
    "Insurance_AR": ("Original Arabic insurance label", "Raw"),
    "Is_Uninsured": ("TRUE when Insurance = Uninsured", "Calculated"),
    "Severity": ("Low / Medium / High (ordered)", "Raw, translated"),
    "Severity_AR": ("Original Arabic severity label", "Raw"),
    "High_Severity": ("TRUE when Severity = High", "Calculated"),
    "Bill": ("Genuine bill amount; blank where the raw value was a 999/3852 placeholder", "Cleaned"),
    "Bill_Raw": ("Bill exactly as in the source file", "Raw"),
    "Bill_Status": ("Recorded / Placeholder 999 / Placeholder 3852", "Calculated"),
    "Bill_Per_Day": ("Bill / Length_of_Stay (genuine bills only)", "Calculated"),
    "Date_Corrected": ("TRUE when admission/discharge dates were reversed and swapped", "Calculated"),
    "Review_Flag": ("'Long stay for minor condition' or 'Atypical age for diagnosis'", "Calculated"),
}
ws_dd = wb.create_sheet("Data Dictionary")
title(ws_dd, "Data dictionary", "Column meaning, type and origin for the Clean Data sheet.")
header(ws_dd, 4, ["Column", "Description", "Type", "Origin"])
for i, col in enumerate(df.columns, 5):
    desc, origin = DICT[col]
    kind = ("Date" if "Date" in col and "Corrected" not in col else "Boolean" if df[col].dtype == bool else
            "Number" if pd.api.types.is_numeric_dtype(df[col]) else "Text")
    for j, v in enumerate([col, desc, kind, origin], 1):
        cell = ws_dd.cell(row=i, column=j, value=v)
        cell.alignment, cell.border = WRAP, BOX
    ws_dd.cell(row=i, column=1).font = F_BOLD
widths(ws_dd, [22, 80, 10, 16])

# ------------------------------------------------------------------ Raw Data (untouched copy)
raw_wb = load_workbook("Hospital Analytics.xlsx")
raw_ws = raw_wb.active
ws_raw = wb.create_sheet("Raw Data")
for r in raw_ws.iter_rows(values_only=True):
    ws_raw.append(list(r))
for cell in ws_raw["F"][1:] + ws_raw["G"][1:]:
    cell.number_format = "yyyy-mm-dd"
ws_raw.sheet_view.rightToLeft = True
widths(ws_raw, [18, 8, 8, 14, 18, 14, 14, 10, 14, 14])
for c in ws_raw[1]:
    c.font, c.fill = F_HEAD, FILL_HEAD

# ------------------------------------------------------------------ About
ws_a = wb.create_sheet("About", 0)
title(ws_a, "Hospital Analytics 2023 - cleaned data & measures", "VOLTIX Data Analyst track - Task 7")
ABOUT = [
    ("KPIs", "19 headline measures (volume, efficiency, acuity, finance, data quality) - each with formula logic and purpose."),
    ("Summary Tables", "Admissions, ALOS, bed-days, long-stay, severity, uninsured and billing measures by 10 dimensions."),
    ("Severity x Dx", "Severity mix per diagnosis - shows the severity field does not line up with diagnosis or stay length."),
    ("Daily Census", "Patients in a bed each night of 2023 - capacity view behind Avg / Peak census."),
    ("Clean Data", f"{N} cleaned admissions (Excel Table 'tblPatients') with {df.shape[1]} columns incl. calculated columns."),
    ("Cleaning Log", "Every check, rows affected and action taken (duplicates, reversed dates, placeholder bills...)."),
    ("Data Dictionary", "Meaning and origin of every column."),
    ("Raw Data", "The original sheet, unchanged, for audit."),
]
header(ws_a, 4, ["Sheet", "Contents"])
for i, (s, t) in enumerate(ABOUT, 5):
    ws_a.cell(row=i, column=1, value=s).font = F_BOLD
    ws_a.cell(row=i, column=1).hyperlink = f"#'{s}'!A1"
    ws_a.cell(row=i, column=2, value=t).alignment = WRAP
r = 5 + len(ABOUT) + 1
for line in [
    "Key cleaning decisions",
    "- 1 exact duplicate removed; 29 reversed admission/discharge dates swapped (flagged in Date_Corrected).",
    "- Bills of 999 (x100) and 3852 (x95) are system placeholders: identical across every diagnosis, severity and stay length,",
    "  while the 52 genuine bills rise with stay length (r = 0.64). They are blanked, not imputed; Bill_Raw keeps the original.",
    "- Unusual but possible values are flagged, not deleted (long stays for minor conditions, atypical ages).",
    "Full dashboard: Hospital_Dashboard.html. Insights: KEY_INSIGHTS.md.",
]:
    ws_a.cell(row=r, column=1, value=line).font = F_SECTION if line == "Key cleaning decisions" else F_BASE
    r += 1
widths(ws_a, [20, 110])

for ws in wb.worksheets:
    base_font(ws)
wb.save(OUT)
print("saved", OUT, [ws.title for ws in wb.worksheets])

# ------------------------------------------------------------------ recalculate (cache formula results)
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
if os.path.exists(SOFFICE):
    tmp = tempfile.mkdtemp()
    subprocess.run([SOFFICE, "--headless", "--norestore", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", tmp, OUT], check=True, capture_output=True, timeout=300)
    shutil.move(os.path.join(tmp, OUT), OUT)
    chk = load_workbook(OUT, data_only=True)
    errors = [f"{ws.title}!{c.coordinate}" for ws in chk for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and len(c.value) > 1]
    print("recalculated with LibreOffice; formula errors:", errors or "none")
else:
    print("LibreOffice not found - open the file in Excel once and save to cache formula results")
